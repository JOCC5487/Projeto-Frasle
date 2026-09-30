from __future__ import annotations
import argparse, csv, hashlib, io, json, logging, mimetypes, re, sqlite3, sys, tempfile, time, zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
import yaml
from openpyxl import load_workbook
from PIL import Image
from xml.etree import ElementTree as ET

NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships','xdr':'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main','pr':'http://schemas.openxmlformats.org/package/2006/relationships'}
OOXML={'.xlsx','.xlsm','.xltx','.xltm'}
LOG=logging.getLogger('dinamometro')

def now(): return datetime.now(timezone.utc).astimezone().isoformat(timespec='seconds')
def sha256_bytes(b): return hashlib.sha256(b).hexdigest()
def sha256_file(p, block=1024*1024):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for c in iter(lambda:f.read(block), b''): h.update(c)
    return h.hexdigest()
def safe(s):
    s=re.sub(r'[<>:"/\\|?*\x00-\x1f]+','_',str(s)).strip(' ._')
    return s[:120] or 'sem_nome'
def qname(tag): return tag.split('}',1)[-1]
def norm_target(base, target):
    from posixpath import normpath, join
    return normpath(join(base, target)).lstrip('/')

def init_db(c):
    c.executescript('''
    PRAGMA journal_mode=WAL; PRAGMA foreign_keys=ON; PRAGMA synchronous=NORMAL;
    CREATE TABLE IF NOT EXISTS execucoes(id INTEGER PRIMARY KEY, inicio TEXT NOT NULL, fim TEXT, origem TEXT, status TEXT, arquivos_avaliados INTEGER DEFAULT 0, processados INTEGER DEFAULT 0, ignorados INTEGER DEFAULT 0, sem_imagens INTEGER DEFAULT 0, com_erro INTEGER DEFAULT 0, imagens_extraidas INTEGER DEFAULT 0, duplicadas INTEGER DEFAULT 0, pendentes INTEGER DEFAULT 0);
    CREATE TABLE IF NOT EXISTS arquivos(id INTEGER PRIMARY KEY, caminho_origem TEXT NOT NULL UNIQUE, arquivo_nome TEXT NOT NULL, extensao TEXT, tamanho_bytes INTEGER, modificado_ns INTEGER, hash_arquivo TEXT, numero_ensaio TEXT, cabecalho_json TEXT, status TEXT, total_imagens INTEGER DEFAULT 0, ultimo_processamento TEXT);
    CREATE TABLE IF NOT EXISTS imagens(id_imagem INTEGER PRIMARY KEY, arquivo_id INTEGER NOT NULL REFERENCES arquivos(id) ON DELETE CASCADE, numero_ensaio TEXT, aba_origem TEXT NOT NULL, etapa_ensaio TEXT, classificacao_imagem TEXT NOT NULL, status_classificacao TEXT NOT NULL, dados_cabecalho TEXT, dados_relacionados TEXT, referencia_imagem TEXT NOT NULL, ancora_celula TEXT, nome_interno TEXT, mime_type TEXT, extensao TEXT, largura_px INTEGER, altura_px INTEGER, tamanho_bytes INTEGER, hash_imagem TEXT NOT NULL, arquivo_imagem TEXT, imagem_blob BLOB, data_processamento TEXT NOT NULL, UNIQUE(arquivo_id, referencia_imagem, hash_imagem));
    CREATE TABLE IF NOT EXISTS erros(id INTEGER PRIMARY KEY, execucao_id INTEGER REFERENCES execucoes(id), caminho_origem TEXT, etapa TEXT, tipo_erro TEXT, mensagem TEXT, data_erro TEXT);
    CREATE TABLE IF NOT EXISTS metadados(chave TEXT PRIMARY KEY, valor TEXT NOT NULL, atualizado_em TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS categorias(codigo TEXT PRIMARY KEY, descricao TEXT NOT NULL, ativa INTEGER NOT NULL DEFAULT 1 CHECK(ativa IN (0,1)));
    INSERT OR IGNORE INTO categorias(codigo,descricao) VALUES
      ('componente','Componente ou corpo de prova'),
      ('avaria','Avaria, falha, trinca, quebra ou desgaste'),
      ('resultado','Resultado, grafico ou medicao'),
      ('montagem','Montagem, instalacao ou bancada'),
      ('identificacao','Etiqueta, codigo, serial ou lote'),
      ('nao_classificada','Imagem aguardando revisao');
    INSERT OR REPLACE INTO metadados(chave,valor,atualizado_em) VALUES('versao_schema','2',datetime('now'));
    CREATE INDEX IF NOT EXISTS ix_arq_nome ON arquivos(arquivo_nome); CREATE INDEX IF NOT EXISTS ix_arq_dl ON arquivos(numero_ensaio); CREATE INDEX IF NOT EXISTS ix_arq_status ON arquivos(status);
    CREATE INDEX IF NOT EXISTS ix_img_arquivo ON imagens(arquivo_id); CREATE INDEX IF NOT EXISTS ix_img_dl ON imagens(numero_ensaio); CREATE INDEX IF NOT EXISTS ix_img_aba ON imagens(aba_origem); CREATE INDEX IF NOT EXISTS ix_img_etapa ON imagens(etapa_ensaio); CREATE INDEX IF NOT EXISTS ix_img_class ON imagens(classificacao_imagem,status_classificacao); CREATE INDEX IF NOT EXISTS ix_img_hash ON imagens(hash_imagem); CREATE INDEX IF NOT EXISTS ix_erro_exec ON erros(execucao_id);
    CREATE VIEW IF NOT EXISTS vw_imagens_consulta AS SELECT i.id_imagem,i.numero_ensaio,a.arquivo_nome,a.caminho_origem,i.aba_origem,i.etapa_ensaio,i.classificacao_imagem,i.status_classificacao,i.ancora_celula,i.referencia_imagem,i.nome_interno,i.mime_type,i.extensao,i.largura_px,i.altura_px,i.tamanho_bytes,i.arquivo_imagem,i.hash_imagem,i.data_processamento FROM imagens i JOIN arquivos a ON a.id=i.arquivo_id;
    CREATE VIEW IF NOT EXISTS vw_resumo_arquivos AS SELECT a.id,a.numero_ensaio,a.arquivo_nome,a.caminho_origem,a.status,a.total_imagens,a.ultimo_processamento,COUNT(i.id_imagem) AS imagens_no_banco,SUM(CASE WHEN i.status_classificacao='pendente' THEN 1 ELSE 0 END) AS pendentes FROM arquivos a LEFT JOIN imagens i ON i.arquivo_id=a.id GROUP BY a.id;
    ''')

def extract_dl(name, text=''):
    pats=[r'(?i)\b(?:DL|DC)[\s_.-]*\d+[A-Z0-9_.-]*',r'(?i)\b(?:DL|DC)[A-Z0-9_.-]+']
    for src in (name,text):
        for p in pats:
            m=re.search(p,src)
            if m: return re.sub(r'\s+','',m.group(0)).upper().strip('._-')
    return None

def workbook_context(path, max_rows, max_cols):
    wb=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    sheets={}; all_header=[]
    for ws in wb.worksheets:
        rows=[]
        for row in ws.iter_rows(min_row=1,max_row=max_rows,max_col=max_cols,values_only=True):
            vals=[str(v).strip() for v in row if v not in (None,'')]
            if vals: rows.append(' | '.join(vals)[:1000])
        sheets[ws.title]=rows
        all_header.extend(f'{ws.title}: {x}' for x in rows[:15])
    wb.close(); return sheets, all_header

def classify(cfg, sheet, nearby, header):
    text=' '.join([sheet,nearby,*header[:20]]).casefold()
    hits=[]
    for cat, words in cfg.get('categorias',{}).items():
        n=sum(1 for w in words if str(w).casefold() in text)
        if n: hits.append((n,cat))
    hits.sort(reverse=True)
    if not hits: return 'nao_classificada','pendente'
    if len(hits)>1 and hits[0][0]==hits[1][0]: return 'nao_classificada','pendente'
    return hits[0][1],'automatica'

def stage(sheet, nearby):
    for token in ['antes','inicial','pre-ensaio','durante','intermediario','depois','apos','final','pos-ensaio','montagem','resultado']:
        if token in f'{sheet} {nearby}'.casefold(): return token
    return sheet

def xlsx_images(path, sheet_rows):
    out=[]
    with zipfile.ZipFile(path) as z:
        wb=ET.fromstring(z.read('xl/workbook.xml'))
        rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        wbrel={e.attrib['Id']:e.attrib['Target'] for e in rels}
        for sh in wb.find('m:sheets',NS):
            title=sh.attrib['name']; rid=sh.attrib.get('{%s}id'%NS['r']); sheet_path=norm_target('xl',wbrel[rid])
            try: root=ET.fromstring(z.read(sheet_path))
            except KeyError: continue
            drawing=root.find('m:drawing',NS)
            if drawing is None: continue
            drid=drawing.attrib.get('{%s}id'%NS['r']); relpath=str(Path(sheet_path).parent).replace('\\','/')+'/_rels/'+Path(sheet_path).name+'.rels'
            srels=ET.fromstring(z.read(relpath)); st={e.attrib['Id']:e.attrib['Target'] for e in srels}; draw_path=norm_target(str(Path(sheet_path).parent).replace('\\','/'),st[drid])
            droot=ET.fromstring(z.read(draw_path)); drels_path=str(Path(draw_path).parent).replace('\\','/')+'/_rels/'+Path(draw_path).name+'.rels'; drels=ET.fromstring(z.read(drels_path)); dt={e.attrib['Id']:e.attrib['Target'] for e in drels}
            idx=0
            for anchor in list(droot):
                blip=anchor.find('.//a:blip',NS)
                if blip is None: continue
                erid=blip.attrib.get('{%s}embed'%NS['r']); target=dt.get(erid)
                if not target: continue
                media=norm_target(str(Path(draw_path).parent).replace('\\','/'),target); data=z.read(media); idx+=1
                fr=anchor.find('xdr:from',NS); col=row=0
                if fr is not None:
                    col=int(fr.findtext('xdr:col','0',NS)); row=int(fr.findtext('xdr:row','0',NS))
                cell=f'{col+1},{row+1}'
                nearby=' | '.join(sheet_rows.get(title,[])[max(0,row-3):row+4])[:3000]
                name=Path(media).name; ref=f'{title}|{qname(anchor.tag)}|{cell}|{idx}|{name}'
                out.append(dict(sheet=title,data=data,name=name,ref=ref,anchor=cell,nearby=nearby))
    return out

def legacy_images_com(path):
    try: import win32com.client
    except ImportError as e: raise RuntimeError('pywin32 ausente para arquivo legado') from e
    app=win32com.client.DispatchEx('Excel.Application'); app.Visible=False; app.DisplayAlerts=False; wb=None; out=[]
    try:
        wb=app.Workbooks.Open(str(path),ReadOnly=True,UpdateLinks=0)
        with tempfile.TemporaryDirectory() as td:
            for ws in wb.Worksheets:
                for i in range(1,ws.Shapes.Count+1):
                    shape=ws.Shapes.Item(i)
                    if shape.Type not in (11,13): continue
                    tmp=Path(td)/f'{safe(ws.Name)}_{i}.png'
                    chart=ws.ChartObjects().Add(0,0,max(1,shape.Width),max(1,shape.Height)); shape.Copy(); chart.Chart.Paste(); chart.Chart.Export(str(tmp),'PNG'); chart.Delete()
                    cell=shape.TopLeftCell.Address(False,False); out.append(dict(sheet=ws.Name,data=tmp.read_bytes(),name=f'{safe(ws.Name)}_{i}.png',ref=f'{ws.Name}|shape|{cell}|{i}',anchor=cell,nearby=''))
        return out
    finally:
        if wb: wb.Close(False)
        app.Quit()

def image_meta(data,name):
    mime=mimetypes.guess_type(name)[0] or 'application/octet-stream'; w=h=None; ext=Path(name).suffix.lower() or '.bin'
    try:
        with Image.open(io.BytesIO(data)) as im:
            w,h=im.size; fmt=(im.format or '').lower(); ext='.'+('jpg' if fmt=='jpeg' else fmt); mime=Image.MIME.get(im.format,mime)
    except Exception: pass
    return mime,ext,w,h

def save_report(outdir, exec_id, summary, details):
    rdir=outdir/'relatorios'; rdir.mkdir(parents=True,exist_ok=True); base=rdir/f'execucao_{exec_id}'
    (base.with_suffix('.json')).write_text(json.dumps({'resumo':summary,'arquivos':details},ensure_ascii=False,indent=2),encoding='utf-8')
    keys=['caminho','status','imagens','mensagem']
    with base.with_suffix('.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=keys,extrasaction='ignore'); w.writeheader(); w.writerows(details)

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--config',default='config.yaml'); args=ap.parse_args()
    cfg=yaml.safe_load(Path(args.config).read_text(encoding='utf-8')); source=Path(cfg['origem']); outdir=Path(cfg['saida']).resolve(); outdir.mkdir(parents=True,exist_ok=True); (outdir/'logs').mkdir(exist_ok=True)
    logging.basicConfig(level=logging.INFO,format='%(asctime)s | %(levelname)s | %(message)s',handlers=[logging.FileHandler(outdir/'logs/processamento.log',encoding='utf-8'),logging.StreamHandler(sys.stdout)])
    if not source.exists(): raise SystemExit(f'Origem nao encontrada: {source}')
    db=Path(cfg['banco']).resolve(); con=sqlite3.connect(db); con.row_factory=sqlite3.Row; init_db(con)
    cur=con.execute('INSERT INTO execucoes(inicio,origem,status) VALUES(?,?,?)',(now(),str(source),'executando')); eid=cur.lastrowid; con.commit()
    stats=dict(arquivos_avaliados=0,processados=0,ignorados=0,sem_imagens=0,com_erro=0,imagens_extraidas=0,duplicadas=0,pendentes=0); details=[]
    name_re=re.compile(cfg.get('filtro_nome_regex','(?i)(DC|DL)')); exts={x.lower() for x in cfg['extensoes']}
    candidates=[]
    for p in source.rglob('*'):
        if p.is_file() and p.suffix.lower() in exts and not p.name.startswith('~$'):
            stats['arquivos_avaliados']+=1
            if name_re.search(p.stem): candidates.append(p)
            else: stats['ignorados']+=1
    for n,p in enumerate(candidates,1):
        LOG.info('[%s/%s] %s',n,len(candidates),p)
        try:
            st=p.stat(); existing=con.execute('SELECT * FROM arquivos WHERE caminho_origem=?',(str(p),)).fetchone()
            if existing and existing['tamanho_bytes']==st.st_size and existing['modificado_ns']==st.st_mtime_ns and existing['status']=='sucesso':
                stats['ignorados']+=1; details.append({'caminho':str(p),'status':'inalterado','imagens':existing['total_imagens'],'mensagem':''}); continue
            filehash=sha256_file(p); sheets={}; headers=[]
            if p.suffix.lower() in OOXML:
                sheets,headers=workbook_context(p,int(cfg.get('limite_cabecalho_linhas',40)),int(cfg.get('limite_cabecalho_colunas',30))); imgs=xlsx_images(p,sheets)
            elif cfg.get('usar_excel_com_para_formatos_legados',True): imgs=legacy_images_com(p)
            else: raise RuntimeError('Formato legado desabilitado')
            header_text='\n'.join(headers); dl=extract_dl(p.stem,header_text); header_json=json.dumps(headers,ensure_ascii=False)
            con.execute('''INSERT INTO arquivos(caminho_origem,arquivo_nome,extensao,tamanho_bytes,modificado_ns,hash_arquivo,numero_ensaio,cabecalho_json,status,total_imagens,ultimo_processamento) VALUES(?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(caminho_origem) DO UPDATE SET arquivo_nome=excluded.arquivo_nome,extensao=excluded.extensao,tamanho_bytes=excluded.tamanho_bytes,modificado_ns=excluded.modificado_ns,hash_arquivo=excluded.hash_arquivo,numero_ensaio=excluded.numero_ensaio,cabecalho_json=excluded.cabecalho_json,status='processando',ultimo_processamento=excluded.ultimo_processamento''',(str(p),p.name,p.suffix.lower(),st.st_size,st.st_mtime_ns,filehash,dl,header_json,'processando',0,now()))
            fid=con.execute('SELECT id FROM arquivos WHERE caminho_origem=?',(str(p),)).fetchone()[0]; seen=[]; inserted=0
            for im in imgs:
                hsh=sha256_bytes(im['data']); seen.append((im['ref'],hsh)); cat,status=classify(cfg,im['sheet'],im['nearby'],headers); et=stage(im['sheet'],im['nearby']); mime,ext,w,h=image_meta(im['data'],im['name'])
                imgpath=None
                if cfg.get('salvar_arquivos_imagem',True):
                    folder=outdir/'imagens'/safe(dl or 'sem_dl')/safe(p.stem); folder.mkdir(parents=True,exist_ok=True); imgpath=folder/f'{safe(im["sheet"])}_{safe(im["anchor"])}_{hsh[:12]}{ext}'; imgpath.write_bytes(im['data'])
                before=con.total_changes
                con.execute('''INSERT OR IGNORE INTO imagens(arquivo_id,numero_ensaio,aba_origem,etapa_ensaio,classificacao_imagem,status_classificacao,dados_cabecalho,dados_relacionados,referencia_imagem,ancora_celula,nome_interno,mime_type,extensao,largura_px,altura_px,tamanho_bytes,hash_imagem,arquivo_imagem,imagem_blob,data_processamento) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',(fid,dl,im['sheet'],et,cat,status,header_json,im['nearby'],im['ref'],im['anchor'],im['name'],mime,ext,w,h,len(im['data']),hsh,str(imgpath) if imgpath else None,sqlite3.Binary(im['data']) if cfg.get('armazenar_blob',True) else None,now()))
                if con.total_changes>before: inserted+=1
                else: stats['duplicadas']+=1
                if status=='pendente': stats['pendentes']+=1
            if seen:
                valid={(r,h) for r,h in seen}; old=con.execute('SELECT id_imagem,referencia_imagem,hash_imagem,arquivo_imagem FROM imagens WHERE arquivo_id=?',(fid,)).fetchall()
                for r in old:
                    if (r['referencia_imagem'],r['hash_imagem']) not in valid:
                        con.execute('DELETE FROM imagens WHERE id_imagem=?',(r['id_imagem'],))
            else: con.execute('DELETE FROM imagens WHERE arquivo_id=?',(fid,))
            con.execute("UPDATE arquivos SET status='sucesso',total_imagens=?,ultimo_processamento=? WHERE id=?",(len(imgs),now(),fid)); con.commit()
            stats['processados']+=1; stats['imagens_extraidas']+=inserted
            if not imgs: stats['sem_imagens']+=1
            details.append({'caminho':str(p),'status':'sucesso' if imgs else 'sem_imagens','imagens':len(imgs),'mensagem':''})
        except Exception as e:
            con.rollback(); stats['com_erro']+=1; LOG.exception('Falha em %s',p); con.execute('INSERT INTO erros(execucao_id,caminho_origem,etapa,tipo_erro,mensagem,data_erro) VALUES(?,?,?,?,?,?)',(eid,str(p),'processamento',type(e).__name__,str(e)[:4000],now())); con.commit(); details.append({'caminho':str(p),'status':'erro','imagens':0,'mensagem':str(e)})
    status='concluido_com_erros' if stats['com_erro'] else 'concluido'
    con.execute('''UPDATE execucoes SET fim=?,status=?,arquivos_avaliados=?,processados=?,ignorados=?,sem_imagens=?,com_erro=?,imagens_extraidas=?,duplicadas=?,pendentes=? WHERE id=?''',(now(),status,*stats.values(),eid)); con.commit(); summary={'execucao_id':eid,'status':status,**stats}; save_report(outdir,eid,summary,details); con.close(); print(json.dumps(summary,ensure_ascii=False,indent=2))
if __name__=='__main__': main()
